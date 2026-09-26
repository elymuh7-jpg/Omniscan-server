"""
OmniScan Server — API de conversion de documents.

Phase 1 : Word (.docx) -> PDF via LibreOffice en mode headless.
Phase 2 : PDF -> Word (.docx) via pdf2docx (texte, tableaux et images
reconstruits par une vraie analyse de mise en page), avec un post-traitement
qui force des bordures visibles sur tous les tableaux — pdf2docx ne
reproduit pas toujours fidèlement les lignes de grille de l'original.

pdf2docx s'appuie sur PyMuPDF, qui n'est pas thread-safe : deux conversions
PDF->Word lancées en même temps dans des threads différents peuvent se
corrompre silencieusement l'une l'autre. FastAPI exécute chaque endpoint
"def" (non "async def") dans un thread séparé, donc deux requêtes qui
arrivent proches dans le temps (ou une requête de conversion qui croise une
requête de monitoring) peuvent se chevaucher. On sérialise donc toutes les
conversions PDF->Word avec un verrou (pdf2docx_lock) : une seule à la fois.

Lancer en local pour tester :
    uvicorn main:app --host 0.0.0.0 --port 8000 --reload
"""

import os
import shutil
import subprocess
import tempfile
import threading
import uuid
from typing import List

import pytesseract
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from fastapi import BackgroundTasks, FastAPI, File, HTTPException, UploadFile
from fastapi.responses import FileResponse
from img2table.document import Image as I2TImage
from img2table.ocr import TesseractOCR
import cv2
import numpy as np
from pdf2docx import Converter
from PIL import Image as PILImage
from rembg import new_session, remove

app = FastAPI(title="OmniScan Server", version="0.3.0")

# Dossier de travail temporaire pour les fichiers reçus/générés.
WORK_DIR = os.path.join(tempfile.gettempdir(), "omniscan_server")
os.makedirs(WORK_DIR, exist_ok=True)

# Durée max qu'on laisse à LibreOffice pour convertir un document (secondes).
# Au-delà, on considère que ça a planté plutôt que de bloquer indéfiniment.
CONVERSION_TIMEOUT_SECONDS = 90

# Verrou global : garantit qu'une seule conversion pdf2docx (PyMuPDF) tourne
# à la fois dans tout le processus, quel que soit le nombre de requêtes
# reçues en parallèle.
pdf2docx_lock = threading.Lock()

# img2table (via OpenCV) n'est pas garanti thread-safe pour des appels
# strictement simultanés — même précaution que pour pdf2docx/PyMuPDF.
table_detection_lock = threading.Lock()

# Instance partagée du moteur OCR Tesseract (français), réutilisée pour
# toutes les requêtes plutôt que recréée à chaque appel.
_tesseract_ocr = TesseractOCR(n_threads=1, lang="fra")

# rembg (U^2-Net) via onnxruntime n'est pas garanti thread-safe pour des
# inférences strictement simultanées — même précaution que pour pdf2docx et
# img2table. Session réutilisée entre requêtes plutôt que recréée à chaque
# appel (son chargement est coûteux).
rembg_lock = threading.Lock()
_rembg_session = new_session("u2net")


def _apply_borders_to_table(table) -> None:
    """
    Force une grille noire complète (haut, bas, gauche, droite, lignes internes)
    sur un objet Table python-docx donné. Factorisé pour être réutilisé aussi
    bien par force_table_borders (tableaux issus de pdf2docx) que par l'endpoint
    de détection de tableaux dans les scans (tableaux construits cellule par
    cellule à partir d'img2table).
    """
    tbl_pr = table._tbl.tblPr

    existing_borders = tbl_pr.find(qn("w:tblBorders"))
    if existing_borders is not None:
        tbl_pr.remove(existing_borders)

    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        edge_element = OxmlElement(f"w:{edge}")
        edge_element.set(qn("w:val"), "single")
        edge_element.set(qn("w:sz"), "4")
        edge_element.set(qn("w:space"), "0")
        edge_element.set(qn("w:color"), "000000")
        borders.append(edge_element)
    tbl_pr.append(borders)


def force_table_borders(docx_path: str) -> None:
    """
    Force des bordures visibles (grille complète) sur tous les tableaux d'un
    document. pdf2docx crée parfois la bonne structure de tableau (lignes,
    colonnes) mais sans reproduire fidèlement les lignes de grille visibles
    de l'original — on les rajoute systématiquement, quitte à ne pas
    correspondre exactement au style d'origine, pour garantir un tableau
    lisible et utilisable.
    """
    document = Document(docx_path)
    print(f"[DEBUG] Nombre de tableaux détectés : {len(document.tables)}")
    if not document.tables:
        return

    for table in document.tables:
        _apply_borders_to_table(table)

    document.save(docx_path)


@app.get("/health")
def health_check():
    """Endpoint simple pour vérifier que le serveur répond (utile pour monitoring)."""
    return {"status": "ok"}


@app.post("/convert/docx-to-pdf")
async def convert_docx_to_pdf(background_tasks: BackgroundTasks, file: UploadFile = File(...)):
    """
    Reçoit un fichier .docx, le convertit en PDF via LibreOffice headless,
    et renvoie le PDF résultant. Préserve fidèlement logos, tableaux et
    mise en page — contrairement à la conversion faite à la main côté app.
    """
    if not file.filename.lower().endswith(".docx"):
        raise HTTPException(status_code=400, detail="Le fichier doit être un .docx")

    request_id = str(uuid.uuid4())
    request_dir = os.path.join(WORK_DIR, request_id)
    os.makedirs(request_dir, exist_ok=True)

    input_path = os.path.join(request_dir, "input.docx")
    expected_output_path = os.path.join(request_dir, "input.pdf")
    lo_profile_dir = os.path.join(request_dir, "lo_profile")

    background_tasks.add_task(shutil.rmtree, request_dir, ignore_errors=True)

    try:
        with open(input_path, "wb") as buffer:
            shutil.copyfileobj(file.file, buffer)

        result = subprocess.run(
            [
                "soffice",
                "--headless",
                "--norestore",
                f"-env:UserInstallation=file://{lo_profile_dir}",
                "--convert-to", "pdf",
                "--outdir", request_dir,
                input_path,
            ],
            capture_output=True,
            text=True,
            timeout=CONVERSION_TIMEOUT_SECONDS,
        )

        if result.returncode != 0 or not os.path.exists(expected_output_path):
            raise HTTPException(
                status_code=500,
                detail=f"Échec de la conversion LibreOffice : {result.stderr.strip()}",
            )

        return FileResponse(
            expected_output_path,
            media_type="application/pdf",
            filename=os.path.splitext(file.filename)[0] + ".pdf",
        )

    except subprocess.TimeoutExpired:
        raise HTTPException(status_code=504, detail="La conversion a dépassé le délai autorisé")


@app.post("/compress/pdf")
def compress_pdf(background_tasks: BackgroundTasks, file: UploadFile = File(...), quality: int = 50):
    """
    Reçoit un fichier .pdf et le recompresse via Ghostscript (ré-encodage des
    images internes, sous-échantillonnage des résolutions, nettoyage de la
    structure) — bien plus efficace qu'une compression "maison", car
    Ghostscript retravaille vraiment le contenu du PDF plutôt que de le
    manipuler en surface.

    [quality] (0-100, reçu du slider de l'appli) est mappé vers les
    préréglages standards de Ghostscript :
        <= 40  -> /screen   (~72 dpi, le plus compact)
        <= 70  -> /ebook    (~150 dpi, bon compromis)
        > 70   -> /printer  (~300 dpi, la meilleure qualité)
    """
    if not file.filename.lower().endswith(".pdf"):
        raise HTTPException(status_code=400, detail="Le fichier doit être un .pdf")

    if quality <= 40:
        gs_preset = "/screen"
    elif quality <= 70:
        gs_preset = "/ebook"
    else:
        gs_preset = "/printer"

    request_id = str(uuid.uuid4())
    request_dir = os.path.join(WORK_DIR, request_id)
    os.makedirs(request_dir, exist_ok=True)

    input_path = os.path.join(request_dir, "input.pdf")
    output_path = os.path.join(request_dir, "compressed.pdf")

    background_tasks.add_task(shutil.rmtree, request_dir, ignore_errors=True)

    try:
        with open(input_path, "wb") as buffer:
            shutil.copyfileobj(file.file, buffer)

        result = subprocess.run(
            [
                "gs",
                "-sDEVICE=pdfwrite",
                "-dCompatibilityLevel=1.4",
                f"-dPDFSETTINGS={gs_preset}",
                "-dNOPAUSE",
                "-dBATCH",
                "-dQUIET",
                f"-sOutputFile={output_path}",
                input_path,
            ],
            capture_output=True,
            text=True,
            timeout=CONVERSION_TIMEOUT_SECONDS,
        )

        if result.returncode != 0 or not os.path.exists(output_path):
            raise HTTPException(
                status_code=500,
                detail=f"Échec de la compression Ghostscript : {result.stderr.strip()}",
            )

        # Filet de sécurité : sur un PDF déjà très optimisé, Ghostscript peut
        # ressortir un fichier légèrement PLUS gros (ré-encodage moins
        # efficace que l'original). Dans ce cas, autant renvoyer l'original.
        if os.path.getsize(output_path) >= os.path.getsize(input_path):
            return FileResponse(
                input_path,
                media_type="application/pdf",
                filename=os.path.splitext(file.filename)[0] + "_compressed.pdf",
            )

        return FileResponse(
            output_path,
            media_type="application/pdf",
            filename=os.path.splitext(file.filename)[0] + "_compressed.pdf",
        )

    except subprocess.TimeoutExpired:
        raise HTTPException(status_code=504, detail="La compression a dépassé le délai autorisé")
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Erreur de compression : {str(e)}")


@app.post("/convert/pdf-to-docx")
def convert_pdf_to_docx(background_tasks: BackgroundTasks, file: UploadFile = File(...)):
    """
    Reçoit un fichier .pdf, le convertit en .docx via pdf2docx, force des
    bordures visibles sur tous les tableaux détectés, et renvoie le résultat.

    Définie en fonction normale (pas "async def") : pdf2docx est bloquant —
    FastAPI l'exécute alors automatiquement dans un thread séparé. La
    conversion elle-même est protégée par pdf2docx_lock pour éviter que
    deux threads touchent PyMuPDF en même temps (voir note en haut du
    fichier).
    """
    if not file.filename.lower().endswith(".pdf"):
        raise HTTPException(status_code=400, detail="Le fichier doit être un .pdf")

    request_id = str(uuid.uuid4())
    request_dir = os.path.join(WORK_DIR, request_id)
    os.makedirs(request_dir, exist_ok=True)

    input_path = os.path.join(request_dir, "input.pdf")
    output_path = os.path.join(request_dir, "output.docx")

    background_tasks.add_task(shutil.rmtree, request_dir, ignore_errors=True)

    try:
        with open(input_path, "wb") as buffer:
            shutil.copyfileobj(file.file, buffer)

        with pdf2docx_lock:
            converter = Converter(input_path)
            try:
                converter.convert(output_path, start=0, end=None)
            finally:
                converter.close()

        if not os.path.exists(output_path):
            raise HTTPException(status_code=500, detail="La conversion PDF vers Word a échoué")

        try:
            force_table_borders(output_path)
        except Exception as e:
            # Un échec du post-traitement des bordures ne doit pas faire perdre
            # toute la conversion : on renvoie quand même le document.
            print(f"Avertissement : échec du forçage des bordures ({e})")

        return FileResponse(
            output_path,
            media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            filename=os.path.splitext(file.filename)[0] + ".docx",
        )

    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Erreur de conversion pdf2docx : {str(e)}")


def _extract_page_content(image_path: str) -> tuple[list, str]:
    """
    Analyse une image de page scannée : détecte les tableaux (via img2table +
    Tesseract, qui repère les lignes de grille visuellement dans l'image —
    contrairement à pdf2docx qui a besoin de données vectorielles absentes
    d'un scan) et, s'il n'y a pas de tableau, fait un simple OCR du texte.

    Renvoie (tables, texte_brut) : "tables" est une liste de DataFrames
    pandas (une par tableau détecté), "texte_brut" n'est rempli que si
    aucun tableau n'a été trouvé sur la page.
    """
    with table_detection_lock:
        doc = I2TImage(src=image_path)
        extracted = doc.extract_tables(
            ocr=_tesseract_ocr,
            implicit_rows=False,
            borderless_tables=True,
        )

    if extracted:
        return [table.df for table in extracted], ""

    # Aucun tableau détecté sur cette page : on se contente d'un OCR classique.
    raw_text = pytesseract.image_to_string(PILImage.open(image_path), lang="fra")
    return [], raw_text.strip()


@app.post("/ocr/images-to-docx")
def ocr_images_to_docx(background_tasks: BackgroundTasks, files: List[UploadFile] = File(...)):
    """
    Reçoit une ou plusieurs images de pages scannées et renvoie un unique
    .docx : chaque page dont un tableau a été détecté devient un vrai
    tableau Word avec bordures (via img2table + Tesseract), chaque page sans
    tableau devient un simple paragraphe de texte OCR. Remplace, pour
    l'export Word du scanner, l'extraction ML Kit locale qui ne récupère que
    du texte brut même quand la page contenait un tableau.
    """
    if not files:
        raise HTTPException(status_code=400, detail="Aucune image reçue")

    request_id = str(uuid.uuid4())
    request_dir = os.path.join(WORK_DIR, request_id)
    os.makedirs(request_dir, exist_ok=True)
    output_path = os.path.join(request_dir, "scan_result.docx")

    background_tasks.add_task(shutil.rmtree, request_dir, ignore_errors=True)

    try:
        document = Document()

        for page_index, upload in enumerate(files, start=1):
            image_path = os.path.join(request_dir, f"page_{page_index}.jpg")
            with open(image_path, "wb") as buffer:
                shutil.copyfileobj(upload.file, buffer)

            if len(files) > 1:
                document.add_heading(f"Page {page_index}", level=2)

            try:
                tables, raw_text = _extract_page_content(image_path)
            except Exception as e:
                print(f"Avertissement : échec de l'analyse de la page {page_index} ({e})")
                tables, raw_text = [], ""

            if tables:
                for df in tables:
                    n_rows, n_cols = df.shape
                    word_table = document.add_table(rows=n_rows + 1, cols=n_cols)
                    _apply_borders_to_table(word_table)

                    for col_index, col_name in enumerate(df.columns):
                        word_table.cell(0, col_index).text = str(col_name)
                    for row_index in range(n_rows):
                        for col_index in range(n_cols):
                            value = df.iat[row_index, col_index]
                            word_table.cell(row_index + 1, col_index).text = "" if value is None else str(value)

                    document.add_paragraph("")
            else:
                document.add_paragraph(raw_text if raw_text else "(Aucun texte détecté sur cette page)")

        document.save(output_path)

        return FileResponse(
            output_path,
            media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            filename="scan_result.docx",
        )

    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Erreur lors de l'analyse du scan : {str(e)}")


def _order_corners(pts: "np.ndarray") -> "np.ndarray":
    """Ordonne 4 points en : haut-gauche, haut-droite, bas-droite, bas-gauche."""
    rect = np.zeros((4, 2), dtype="float32")
    s = pts.sum(axis=1)
    rect[0] = pts[np.argmin(s)]
    rect[2] = pts[np.argmax(s)]
    diff = np.diff(pts, axis=1)
    rect[1] = pts[np.argmin(diff)]
    rect[3] = pts[np.argmax(diff)]
    return rect


def _warp_perspective(image: "np.ndarray", pts: "np.ndarray") -> "np.ndarray":
    rect = _order_corners(pts)
    (tl, tr, br, bl) = rect

    width_a = np.linalg.norm(br - bl)
    width_b = np.linalg.norm(tr - tl)
    max_width = max(int(width_a), int(width_b), 100)

    height_a = np.linalg.norm(tr - br)
    height_b = np.linalg.norm(tl - bl)
    max_height = max(int(height_a), int(height_b), 100)

    dst = np.array(
        [[0, 0], [max_width - 1, 0], [max_width - 1, max_height - 1], [0, max_height - 1]],
        dtype="float32",
    )
    matrix = cv2.getPerspectiveTransform(rect, dst)
    return cv2.warpPerspective(image, matrix, (max_width, max_height))


@app.post("/scan/segment-crop")
def segment_crop(file: UploadFile = File(...)):
    """
    Reçoit une photo brute d'un document et renvoie une version recadrée et
    redressée, en isolant d'abord le document du fond via rembg (U^2-Net)
    avant de chercher son contour — bien plus robuste que la détection de
    contours par seuillage (Canny) utilisée localement dans l'app, en
    particulier sur fond de faible contraste ou avec un doigt visible.
    Si aucun contour fiable n'est trouvé, renvoie l'image d'origine
    inchangée plutôt que d'échouer.
    """
    if not any(file.filename.lower().endswith(ext) for ext in (".jpg", ".jpeg", ".png")):
        raise HTTPException(status_code=400, detail="Le fichier doit être une image (.jpg, .jpeg ou .png)")

    request_id = str(uuid.uuid4())
    request_dir = os.path.join(WORK_DIR, request_id)
    os.makedirs(request_dir, exist_ok=True)
    input_path = os.path.join(request_dir, "input.jpg")
    output_path = os.path.join(request_dir, "cropped.jpg")

    try:
        with open(input_path, "wb") as buffer:
            shutil.copyfileobj(file.file, buffer)

        original = PILImage.open(input_path).convert("RGB")

        with rembg_lock:
            removed = remove(original, session=_rembg_session)

        alpha = np.array(removed)[:, :, 3]
        _, mask = cv2.threshold(alpha, 10, 255, cv2.THRESH_BINARY)
        contours, _ = cv2.findContours(mask, cv2.RETR_EXTERNAL, cv2.CHAIN_APPROX_SIMPLE)

        src_np = np.array(original)

        if not contours:
            # rembg n'a rien isolé de façon fiable : on garde l'image telle quelle
            # plutôt que de risquer un mauvais recadrage.
            PILImage.fromarray(src_np).save(output_path, "JPEG", quality=95)
        else:
            largest = max(contours, key=cv2.contourArea)
            peri = cv2.arcLength(largest, True)
            approx = cv2.approxPolyDP(largest, 0.02 * peri, True)

            if len(approx) == 4:
                pts = approx.reshape(4, 2).astype("float32")
                warped = _warp_perspective(src_np, pts)
            else:
                # Contour fiable mais pas franchement quadrilatère : un simple
                # rectangle englobant reste plus sûr qu'une perspective forcée
                # sur des points mal définis.
                x, y, w, h = cv2.boundingRect(largest)
                warped = src_np[y:y + h, x:x + w]

            PILImage.fromarray(warped).save(output_path, "JPEG", quality=95)

        return FileResponse(output_path, media_type="image/jpeg", filename="cropped.jpg")

    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Erreur lors du recadrage : {str(e)}")
    finally:
        shutil.rmtree(request_dir, ignore_errors=True)
